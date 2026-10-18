"""Interviews govern execution; evidence, privacy, costs and owner gates remain real."""
import json,unittest
from pathlib import Path
from unittest.mock import patch
from decimal import Decimal
from test_milestones import MilestoneBase
from test_assurance import AssuranceBase
from superhuman.projects import engine as p
from superhuman.team import store as team,work,budget
from superhuman.milestones import store as ms,workflow as mw
from superhuman.practice import store as s,web,impact,billing
from superhuman.practice.policy import questions,DEFERRED
from superhuman.navigation import shell

class PracticeTests(MilestoneBase):
 def setUp(self):super().setUp();s.setup()
 def brief(self,pid):
  q=team.project(pid);r=s.interview(pid);s.save(pid,r['version'],{k:('answered','Owner '+k) for k in ('outcome','outputs','deadline','success')},q['sensitivity']);s.confirm(pid,s.interview(pid)['version'])
 def project_ready(self):
  pid,mid=self.project();self.brief(pid);return pid,mid
 def test_new_projects_enrolled_after_setup(self):
  pid,mid=self.project();self.assertIsNotNone(s.interview(pid));self.assertFalse(s.current(pid))
 def test_four_essential_answers_required(self):
  pid,mid=self.project()
  with self.assertRaisesRegex(ValueError,'essential'):s.confirm(pid,1)
 def test_unknown_values_are_not_confirmed_as_facts(self):
  pid,mid=self.project();s.save(pid,1,{'outcome':('research','Find it')},'internal')
  with self.assertRaises(ValueError):s.confirm(pid,2)
 def test_short_brief_allows_deeper_unknowns(self):
  pid,mid=self.project_ready();self.assertTrue(s.current(pid));self.assertNotIn('domain_goal',{a['key'] for a in s.answers(pid)})
 def test_unconfirmed_context_is_not_exposed(self):
  pid,mid=self.project();s.save(pid,1,{'current':('answered','PRIVATE SECRET')},'internal');self.assertNotIn('PRIVATE SECRET',json.dumps(s.context(pid,'marketing','internal')))
 def test_stale_form_cannot_overwrite(self):
  pid,mid=self.project();s.save(pid,1,{'current':('answered','Latest')},'internal')
  with self.assertRaises(ValueError):s.save(pid,1,{'current':('answered','Old')},'internal')
  self.assertEqual(s.answers(pid)[0]['content'],'Latest')
 def test_edit_holds_previously_confirmed_work(self):
  pid,mid=self.project_ready();self.start(mid);a=self.lead(mid);self.assertTrue(work.gate(team.task(a['task_id']))[0]);s.save(pid,s.interview(pid)['version'],{'current':('answered','Changed')},'internal');self.assertFalse(work.gate(team.task(a['task_id']))[0])
 def test_unconfirmed_start_blocked(self):
  pid,mid=self.project()
  with self.assertRaisesRegex(ValueError,'interview'):self.start(mid)
 def test_control_change_requires_reconfirmation(self):
  pid,mid=self.project_ready()
  with p.connect() as db:db.execute('UPDATE team_projects SET version=version+1 WHERE id=?',(pid,))
  self.assertFalse(s.current(pid));self.assertTrue(s.hold(pid))
 def test_private_brief_cannot_enter_lower_assignment(self):
  pid,mid=self.project_ready();s.save(pid,s.interview(pid)['version'],{},'confidential')
  with self.assertRaisesRegex(ValueError,'sensitivity'):s.confirm(pid,s.interview(pid)['version'])
 def test_outside_roster_has_no_answers(self):
  pid,mid=self.project_ready();self.assertNotIn('answers',s.context(pid,'legal_ops','restricted'))
 def test_context_contains_confirmed_rules_and_answers(self):
  pid,mid=self.project_ready();ctx=s.context(pid,'marketing','internal');self.assertIn('answers',ctx);self.assertIn('authority',ctx['operating_rules'])
 def test_draft_questionnaire_contains_native_labels_and_steps(self):
  pid,mid=self.project();body=web.interview_page(pid).body.decode();self.assertIn('Interview progress',body);self.assertIn('for="answer-outcome"',body);self.assertIn('Save and continue',body)
 def test_user_content_is_escaped(self):
  pid,mid=self.project();s.save(pid,1,{'outcome':('answered','<script>bad()</script>')},'internal');self.assertNotIn('<script>bad()',web.interview_page(pid,5).body.decode())
 def test_invalid_step_rejected(self):
  pid,mid=self.project()
  with self.assertRaises(ValueError):web.interview_page(pid,20)
 def test_deferred_choices_retained(self):self.assertEqual([q['question'] for q in DEFERRED],list(range(61,68)))
 def test_endorsement_preset_preserves_existing_answer(self):
  pid,mid=self.project(scope='endorsement');s.save(pid,1,{'outcome':('answered','My exact goal')},'internal');s.preset(pid,2);a={r['key']:r for r in s.answers(pid)};self.assertEqual(a['outcome']['content'],'My exact goal');self.assertEqual(a['cost']['status'],'undecided');self.assertIn('workbook',a['constraints']['content'])
 def test_wrong_department_preset_rejected(self):
  pid,mid=self.project()
  with self.assertRaises(ValueError):s.preset(pid,1)
 def test_history_is_append_only(self):
  pid,mid=self.project_ready();rows=s.read('practice_history','SELECT * FROM practice_history WHERE project_id=?',(pid,));self.assertEqual([r['action'] for r in rows],['saved','confirmed'])
 def test_signature_changes_with_answer(self):
  pid,mid=self.project_ready();old=s.signature(pid);s.save(pid,s.interview(pid)['version'],{'current':('answered','New')},'internal');self.assertNotEqual(old,s.signature(pid))
 def test_dispatch_transaction_rechecks_brief(self):
  pid,mid=self.project_ready();self.start(mid);t=team.task(self.lead(mid)['task_id']);s.save(pid,s.interview(pid)['version'],{},'internal')
  with p.connect() as db:
   with self.assertRaisesRegex(ValueError,'changed'):s.dispatch_check(db,t)
 def test_research_derived_questions_scoped_to_actual_task(self):
  pid,mid=self.project_ready();self.start(mid);t=team.task(self.lead(mid)['task_id']);rid=work.dispatch(t['id'],t['version']);qid=s.propose({'project_id':rid},'What is tested?','Clarify product reality');self.assertEqual(qid,s.propose({'project_id':rid},'What is tested?','Duplicate'));self.assertEqual(s.snapshot(pid)['questions'][0]['project_id'],pid)
 def test_pending_questions_appear_once_in_shared_queue(self):
  from superhuman.practice.attention import items
  from superhuman.workspace import model
  pid,mid=self.project_ready();self.start(mid);t=team.task(self.lead(mid)['task_id']);rid=work.dispatch(t['id'],t['version']);s.propose({'project_id':rid},'What is tested?','Clarify')
  self.assertEqual(len([r for r in items() if r['reason']=='Project interview clarification']),1)
  self.assertEqual(len([r for r in model.decisions() if r['reason']=='Project interview clarification']),1)
 def test_proposed_question_does_not_change_confirmed_scope_fingerprint(self):
  pid,mid=self.project_ready();self.start(mid);t=team.task(self.lead(mid)['task_id']);rid=work.dispatch(t['id'],t['version']);before=s.signature(pid);s.propose({'project_id':rid},'What is tested?','Clarify');self.assertEqual(before,s.signature(pid));self.assertTrue(s.current(pid))
 def test_arbitrary_run_cannot_propose(self):
  with self.assertRaises(ValueError):s.propose({'project_id':'unknown'},'Q','Reason')
 def test_clarification_invalidates_confirmation(self):
  pid,mid=self.project_ready();self.start(mid);t=team.task(self.lead(mid)['task_id']);rid=work.dispatch(t['id'],t['version']);qid=s.propose({'project_id':rid},'What is tested?','Clarify');s.answer_question(pid,qid,1,'Only prototype','answered');self.assertFalse(s.current(pid))
 def test_auto_start_opt_in_required(self):
  pid,mid=self.project_ready();self.assertIsNone(s.auto_start());self.assertEqual(ms.milestone(mid)['state'],'planned')
 def test_auto_start_preserves_owner_dependencies(self):
  pid,mid=self.project_ready();s.save(pid,s.interview(pid)['version'],{},'internal',True);s.confirm(pid,s.interview(pid)['version']);later=ms.add(pid,'Dependent','Q','Criteria',['marketing'],deps=[mid]);ms.confirm_criteria(later,ms.milestone(later)['version']);out=s.auto_start();self.assertEqual(out['status'],'milestone_started');self.assertEqual(ms.milestone(mid)['state'],'working');self.assertEqual(ms.milestone(later)['state'],'planned')
 def test_auto_start_held_when_criteria_unconfirmed(self):
  pid,mid=self.project(confirmed=False);self.brief(pid);s.save(pid,s.interview(pid)['version'],{},'internal',True);s.confirm(pid,s.interview(pid)['version']);self.assertIsNone(s.auto_start())
 def test_asset_names_allowlisted(self):
  self.assertEqual(web.portrait('bad'),'');self.assertIn('/assets/agents/twin.png',web.portrait('twin'))
 def test_portraits_all_exist(self):
  for name in team.EMPLOYEES:self.assertTrue((web.ASSETS/(name+'.png')).read_bytes().startswith(b'\x89PNG'))
 def test_unknown_cost_not_inferred_as_roi(self):
  pid,mid=self.project_ready();body=web.impact_page(pid).body.decode();self.assertIn('not a complete project cost or ROI',body);self.assertNotIn('ROI: 100',body)
 def test_nonfinite_costs_rejected(self):
  for v in ('nan','inf','-1','10000001'):
   with self.assertRaises(ValueError):impact.money(v)
 def test_billing_links_within_existing_reservation(self):
  pid,mid=self.project_ready()
  with billing.scope(pid,mid),budget.scope('twin'):ident=budget.reserve('gpt-4o-mini',{'messages':[{'role':'user','content':'Synthetic'}],'max_completion_tokens':10},'chat')
  self.assertEqual(s.read('practice_cost_links','SELECT * FROM practice_cost_links WHERE cost_id=?',(ident,))[0]['project_id'],pid);self.assertEqual(impact.costs(pid)['entries'],1)
 def test_legacy_project_not_silently_enrolled(self):
  with p.connect() as db:db.execute('DROP TABLE practice_interviews')
  pid,mid=self.project();self.assertIsNone(s.interview(pid));self.assertFalse(s.hold(pid))
 def test_rollback_preserves_completed_run_state(self):
  pid,mid=self.project_ready();self.start(mid);tid=self.lead(mid)['task_id'];completed=self.complete(tid)
  from superhuman.agentic import store as agent
  self.assertEqual(agent.get(completed['run_id'])['state'],'completed');s.rollback_permissions();self.assertEqual(agent.get(completed['run_id'])['state'],'completed')
 def test_rollback_retires_tool(self):
  self.assertTrue(any('project.interview' in json.loads(e['allowed_tools']) for e in team.employees()));s.rollback_permissions();self.assertFalse(any('project.interview' in json.loads(e['allowed_tools']) for e in team.employees()))

class ImpactTests(AssuranceBase):
 def setUp(self):super().setUp();s.setup()
 def test_packet_linked_review_and_historical_version(self):
  pid,mid=self.project();r=s.interview(pid);s.save(pid,r['version'],{k:('answered',k) for k in ('outcome','outputs','deadline','success')},'internal');s.confirm(pid,s.interview(pid)['version']);self.confirm(mid);self.start(mid);self.packet(mid);self.accepted(mid);packet=ms.packet(mid)
  ident=impact.add(mid,packet['id'],{k:'not_assessed' for k in impact.CHECKS},'Draft ready','Not observed yet','Independent method deferred',{})
  row=s.read('practice_impacts','SELECT * FROM practice_impacts WHERE id=?',(ident,))[0];self.assertEqual(row['packet_digest'],packet['digest'])
  with self.assertRaises(ValueError):impact.add(mid,'old',{k:'met' for k in impact.CHECKS},'Ready','Done','No',{})

class PracticeHTTPTests(AssuranceBase):
 def setUp(self):
  super().setUp();s.setup()
  import re
  from fastapi.testclient import TestClient
  from superhuman.hosted import config,web,auth
  self.cfg=config.Settings('https://superhuman.example',state_dir=Path('state'));auth.Store(self.cfg.auth_db).setup('owner','synthetic-test-passphrase');self.client=TestClient(web.create_app(self.cfg),base_url=self.cfg.public_origin);self.re=re
  csrf=self.token(self.client.get('/login').text);self.client.post('/login',data=dict(username='owner',password='synthetic-test-passphrase',csrf=csrf))
 def token(self,body):return self.re.search(r'name="csrf" value="([^"]+)"',body).group(1)
 def test_creation_redirects_to_interview(self):
  csrf=self.token(self.client.get('/new-project').text);r=self.client.post('/workspace-projects/create',data=dict(csrf=csrf,scope='marketing',title='New project',objective='Research',acceptance='Useful'),follow_redirects=False);self.assertEqual(r.status_code,303);self.assertTrue(r.headers['location'].startswith('/project-briefs/'));self.assertEqual(self.client.get(r.headers['location']).status_code,200)
 def test_save_confirm_and_read_all_steps(self):
  pid,mid=self.project();csrf=self.token(self.client.get('/project-briefs/'+pid).text);payload=dict(csrf=csrf,expected=1,step=1)
  for key in ('outcome','outputs','deadline','success','constraints'):payload['status_'+key]='answered';payload['answer_'+key]='Owner confirmed '+key
  r=self.client.post('/project-briefs/'+pid+'/answers',data=payload,follow_redirects=False);self.assertEqual(r.status_code,303);self.assertIn('step=2',r.headers['location'])
  for step in range(1,6):self.assertEqual(self.client.get('/project-briefs/'+pid+'?step='+str(step)).status_code,200)
  body=self.client.get('/project-briefs/'+pid+'?step=5').text;csrf=self.token(body);self.assertEqual(self.client.post('/project-briefs/'+pid+'/confirm',data=dict(csrf=csrf,expected=2,confirmed='yes'),follow_redirects=False).status_code,303);self.assertTrue(s.current(pid))
 def test_post_without_csrf_rejected(self):
  pid,mid=self.project();r=self.client.post('/project-briefs/'+pid+'/controls',data=dict(expected=1,sensitivity='public'));self.assertEqual(r.status_code,403);self.assertEqual(s.interview(pid)['version'],1)
 def test_assets_and_guide_remain_authenticated(self):
  from fastapi.testclient import TestClient
  from superhuman.hosted import web
  client=TestClient(web.create_app(self.cfg),base_url=self.cfg.public_origin)
  for url in ('/assets/practice.css','/assets/agents/twin.png','/operating-guide'):
   r=client.get(url,follow_redirects=False);self.assertIn(r.status_code,(303,401))
 def test_portrait_traversal_and_unknown_name_fail(self):
  self.assertEqual(self.client.get('/assets/agents/unknown.png').status_code,404);self.assertEqual(self.client.get('/assets/agents/twin.png').headers['content-type'],'image/png')
 def test_overview_projects_team_settings_and_guide_render(self):
  pid,mid=self.project()
  for url in ('/','/projects','/team','/settings','/operating-guide','/team/projects/'+pid,'/project-briefs/'+pid+'/impact'):
   response=self.client.get(url);self.assertEqual(response.status_code,200,response.text[:300])
  self.assertIn('illustrated AI employee',self.client.get('/team').text);self.assertIn('/assets/practice.css?v=134',self.client.get('/').text);self.assertIn('/assets/practice.js?v=134',self.client.get('/').text);self.assertEqual(self.client.get('/assets/practice.js').status_code,200)
 def test_reading_interview_does_not_write_business_records(self):
  pid,mid=self.project()
  with p.connect() as db:before=[tuple(r) for r in db.execute('SELECT * FROM practice_interviews')]
  for i in range(1,6):self.client.get('/project-briefs/'+pid+'?step='+str(i))
  with p.connect() as db:self.assertEqual(before,[tuple(r) for r in db.execute('SELECT * FROM practice_interviews')])
 def test_invalid_form_shows_recovery_without_success(self):
  pid,mid=self.project();csrf=self.token(self.client.get('/project-briefs/'+pid).text);r=self.client.post('/project-briefs/'+pid+'/confirm',data=dict(csrf=csrf,expected=1,confirmed='yes'),headers={'accept':'text/html'});self.assertEqual(r.status_code,409);self.assertIn('essential',r.text);self.assertFalse(s.current(pid))

class ProductionExportTests(unittest.TestCase):
 def setUp(self):
  import tempfile
  self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
 def tearDown(self):self.temp.cleanup()
 def test_native_word_table_and_source_values(self):
  from superhuman.capabilities import exports
  from docx import Document
  filename=self.root/'test.docx';exports.make_docx(filename,'Evidence package','# Findings\n\nA **supported** claim.\n\n| Source | Finding |\n| --- | --- |\n| Official guide | Verify 42 |\n\n- Next action')
  doc=Document(filename);self.assertEqual(len(doc.tables),1);self.assertEqual(doc.tables[0].rows[1].cells[1].text,'Verify 42');self.assertTrue(any('supported' in p.text for p in doc.paragraphs))
 def test_pdf_contains_table_cells_and_literal_markup_safe(self):
  from superhuman.capabilities import exports
  from pypdf import PdfReader
  filename=self.root/'test.pdf';exports.make_pdf(filename,'Evidence package','# Findings\n\nLiteral <tag> & value.\n\n| Source | Finding |\n| --- | --- |\n| Official guide | Verify 42 |')
  text='\n'.join(p.extract_text() for p in PdfReader(filename).pages);self.assertIn('Verify 42',text);self.assertIn('Literal <tag> & value',text);self.assertIn('Page 1',text)
 def test_unsupported_table_keeps_every_cell(self):
  from superhuman.practice.exports import blocks
  source='| A | B |\n| --- | --- |\n| x | y | z |';result=list(blocks(source));self.assertEqual(result[0][0],'verbatim');self.assertIn('x | y | z',result[0][1])
 def test_many_page_report_keeps_final_marker(self):
  from superhuman.capabilities import exports
  from pypdf import PdfReader
  filename=self.root/'long.pdf';text='\n\n'.join('Paragraph '+str(i)+' evidence applies to this project. '*5 for i in range(100))+'\n\nFINAL EVIDENCE MARKER';exports.make_pdf(filename,'Long report',text);reader=PdfReader(filename);self.assertGreater(len(reader.pages),2);self.assertIn('FINAL EVIDENCE MARKER',reader.pages[-1].extract_text())
