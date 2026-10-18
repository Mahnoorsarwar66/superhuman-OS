"""Deterministic document layout: headings, native tables, bullets and page numbers.

No research, claims of independent review, remote assets or automatic publication.
"""
import re
from xml.sax.saxutils import escape

def blocks(text):
 lines=text.splitlines();i=0;pending=[]
 def paragraph():
  if pending:
   value=' '.join(pending);pending.clear();return ('paragraph',value)
 while i<len(lines):
  line=lines[i]
  # A Markdown table needs a header and delimiter row. Pipe-free prose is never inferred as a table.
  if '|' in line and i+1<len(lines) and re.fullmatch(r'\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*',lines[i+1]):
   item=paragraph()
   if item:yield item
   table=[]
   for row in [line]:table.append([v.strip() for v in row.strip().strip('|').split('|')])
   i+=2
   while i<len(lines) and '|' in lines[i] and lines[i].strip():table.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
   if len(table[0])>6 or any(len(row)!=len(table[0]) for row in table):
    # Preserve unsupported structures as verbatim text, rather than silently dropping cells.
    yield 'verbatim','\n'.join(' | '.join(row) for row in table)
   else:yield 'table',table
   continue
  m=re.match(r'^(#{1,3})\s+(.+)',line)
  bullet=re.match(r'^\s*[-*]\s+(.+)',line)
  if not line.strip() or m or bullet or re.match(r'^\d+[.)]\s',line):
   item=paragraph()
   if item:yield item
   if m:yield 'h'+str(len(m[1])),m[2]
   elif bullet:yield 'bullet',bullet[1]
   elif line.strip():yield 'paragraph',line
  else:pending.append(line)
  i+=1
 item=paragraph()
 if item:yield item

def rich(paragraph,text):
 for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
  run=paragraph.add_run(part);run.bold=bool(i%2)

def docx(path,title,text):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 doc=Document();section=doc.sections[0];section.top_margin=section.bottom_margin=Inches(.8);section.left_margin=section.right_margin=Inches(.85)
 normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor.from_string('243B33');normal.paragraph_format.space_after=Pt(9);normal.paragraph_format.line_spacing=1.16
 for name,size in [('Title',27),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
  style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('28523E');style.paragraph_format.keep_with_next=True
 doc.add_paragraph(title,style='Title')
 for kind,value in blocks(text):
  if kind=='table':
   table=doc.add_table(rows=1,cols=len(value[0]));table.style='Light Shading Accent 1'
   for j,v in enumerate(value[0]):rich(table.rows[0].cells[j].paragraphs[0],v)
   repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
   for row in value[1:]:
    cells=table.add_row().cells
    for j,v in enumerate(row):rich(cells[j].paragraphs[0],v)
   for cell in table.rows[0].cells:
    shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EFE5');cell._tc.get_or_add_tcPr().append(shading)
   doc.add_paragraph()
  else:
   paragraph=doc.add_paragraph(style='Heading '+kind[1] if kind.startswith('h') else 'List Bullet' if kind=='bullet' else 'Normal');rich(paragraph,value)
 footer=section.footer.paragraphs[0];footer.alignment=2;run=footer.add_run('Page ');run.font.size=Pt(9);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.core_properties.title=title;doc.save(path)
 if Document(path).paragraphs[0].text!=title:raise ValueError('DOCX validation failed')
 return ['Native headings, bullets and up to six-column Markdown tables. Images, nested/complex Markdown and complex scripts require visual review.']

def pdf(path,title,text):
 from reportlab.lib import colors
 from reportlab.lib.pagesizes import A4
 from reportlab.lib.styles import ParagraphStyle
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle
 from superhuman.capabilities.exports import pdf_fonts
 normal,bold=pdf_fonts(title+text)
 body=ParagraphStyle('PracticeBody',fontName=normal,fontSize=10.5,leading=16,spaceAfter=9,textColor=colors.HexColor('#243b33'),splitLongWords=True)
 title_style=ParagraphStyle('PracticeTitle',parent=body,fontName=bold,fontSize=25,leading=31,spaceAfter=23,textColor=colors.HexColor('#28523e'),keepWithNext=True)
 headings={str(i):ParagraphStyle('PracticeH'+str(i),parent=body,fontName=bold,fontSize={1:17,2:14,3:12}[i],leading={1:23,2:20,3:18}[i],spaceBefore=15,spaceAfter=8,keepWithNext=True,textColor=colors.HexColor('#28523e')) for i in (1,2,3)}
 cell=ParagraphStyle('PracticeCell',parent=body,fontSize=9,leading=13,spaceAfter=0)
 def markup(value):return re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',escape(value)).replace('\n','<br/>')
 story=[Paragraph(escape(title),title_style)]
 for kind,value in blocks(text):
  if kind=='table':
   data=[[Paragraph(markup(v),cell) for v in row] for row in value];width=(A4[0]-108)/len(value[0]);table=LongTable(data,colWidths=[width]*len(value[0]),repeatRows=1,hAlign='LEFT',splitInRow=1)
   table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8efe5')),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.7,colors.HexColor('#aabcaa')),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#dce5df')),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]));story+=[table,Spacer(1,12)]
  else:story.append(Paragraph(markup(value),headings[kind[1]] if kind.startswith('h') else body,bulletText='•' if kind=='bullet' else None))
 def footer(canvas,doc):
  canvas.saveState();canvas.setStrokeColor(colors.HexColor('#dce5df'));canvas.line(54,43,A4[0]-54,43);canvas.setFont(normal,8);canvas.setFillColor(colors.HexColor('#5d7163'));canvas.drawRightString(A4[0]-54,28,'Page '+str(doc.page));canvas.restoreState()
 SimpleDocTemplate(str(path),pagesize=A4,rightMargin=54,leftMargin=54,topMargin=52,bottomMargin=58,title=title).build(story,onFirstPage=footer,onLaterPages=footer)
 if not path.read_bytes().startswith(b'%PDF-'):raise ValueError('PDF validation failed')
 return ['Native headings, bullets and up to six-column Markdown tables. Images, complex scripts and final pagination require visual review.']
